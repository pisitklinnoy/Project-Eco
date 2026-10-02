"""
Water Level Excel Logger Module
ระบบบันทึกผลการตรวจวัดระดับน้ำลงไฟล์ Excel (water_levels.xlsx) อัตโนมัติ
แยกเก็บเป็นรายคอลัมน์ของแต่ละสถานี โดยบันทึกเฉพาะค่าระดับน้ำ (m R.T.K.) ไม่ใส่ค่า debug
"""

import os
import re
from datetime import datetime
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class WaterLevelExcelLogger:
    """
    คลาสจัดการการบันทึกระดับน้ำลง Excel แบบแยกคอลัมน์รายสถานี
    """
    STATION_COLUMN_NAMES = {
        "muangkong": "Muangkong (m R.T.K.)",
        "x.173a": "Muangkong (m R.T.K.)",
        "bangsala": "Bangsala_X90 (m R.T.K.)",
        "bangsala_x90": "Bangsala_X90 (m R.T.K.)",
        "x.90": "Bangsala_X90 (m R.T.K.)",
        "hatyainai": "Hatyainai_X44 (m R.T.K.)",
        "hatyainai_x44": "Hatyainai_X44 (m R.T.K.)",
        "x.44": "Hatyainai_X44 (m R.T.K.)",
        "station3": "Hatyainai_X44 (m R.T.K.)",
        "ta200304": "Hatyainai_X44 (m R.T.K.)"
    }

    def __init__(self, excel_path: str = None):
        if excel_path is None:
            # ค่าเริ่มต้นบันทึกที่โฟลเดอร์เดียวกับโปรเจกต์ non_time_series/water_levels.xlsx
            curr_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            excel_path = os.path.join(curr_dir, "water_levels.xlsx")
        self.excel_path = os.path.abspath(excel_path)
        self._ensure_file_exists()

    def _normalize_station_column(self, station_name: str) -> str:
        key = str(station_name).strip().lower()
        if key in self.STATION_COLUMN_NAMES:
            return self.STATION_COLUMN_NAMES[key]
        return f"{station_name} (m R.T.K.)"

    def _ensure_file_exists(self):
        """ตรวจสอบและสร้างไฟล์ Excel พร้อมโครงสร้างคอลัมน์มาตรฐาน"""
        if not HAS_OPENPYXL:
            return
        if not os.path.exists(self.excel_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Water Levels"
            
            headers = [
                "Timestamp",
                "Muangkong (m R.T.K.)",
                "Bangsala_X90 (m R.T.K.)",
                "Hatyainai_X44 (m R.T.K.)"
            ]
            ws.append(headers)
            self._apply_header_styling(ws)
            wb.save(self.excel_path)

    def _apply_header_styling(self, ws):
        """ตกแต่งส่วนหัวตาราง Header ให้สวยงามตามสไตล์มืออาชีพ"""
        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='medium', color='1B365D'),
            bottom=Side(style='medium', color='1B365D')
        )

        ws.row_dimensions[1].height = 28
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 25

    def log_water_level(self, station_name: str, water_level_m: float, timestamp: str = None) -> bool:
        """
        บันทึกระดับน้ำลงไฟล์ Excel:
        - station_name: ชื่อหรือรหัสสถานี (เช่น 'Hatyainai_X44', 'X.90', 'Muangkong')
        - water_level_m: ค่าระดับน้ำ (เมตร รทก.)
        - timestamp: วัน-เวลา (เช่น '2026-10-01 12:25:36') ถ้าไม่ระบุจะใช้วันเวลาปัจจุบัน
        """
        if not HAS_OPENPYXL:
            return False

        self._ensure_file_exists()

        if timestamp is None:
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        col_header = self._normalize_station_column(station_name)
        water_level_m = round(float(water_level_m), 2)

        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active

        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]

        # ค้นหาคอลัมน์ของสถานี
        if col_header in headers:
            target_col = headers.index(col_header) + 1
        else:
            target_col = ws.max_column + 1
            ws.cell(row=1, column=target_col, value=col_header)
            self._apply_header_styling(ws)

        # ค้นหาแถวของ Timestamp
        target_row = None
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=1).value
            if val is not None and str(val).strip() == str(timestamp).strip():
                target_row = r
                break

        if target_row is None:
            target_row = ws.max_row + 1
            ws.cell(row=target_row, column=1, value=str(timestamp))

        # บันทึกค่าระดับน้ำ
        cell = ws.cell(row=target_row, column=target_col, value=water_level_m)
        cell.number_format = '0.00'

        # จัดรูปแบบเซลล์
        center_align = Alignment(horizontal="center", vertical="center")
        cell.alignment = center_align

        ws.cell(row=target_row, column=1).alignment = center_align
        ws.row_dimensions[target_row].height = 22

        # เส้นขอบและสีสลับแถว
        thin_border = Border(
            left=Side(style='thin', color='EEEEEE'),
            right=Side(style='thin', color='EEEEEE'),
            top=Side(style='thin', color='EEEEEE'),
            bottom=Side(style='thin', color='EEEEEE')
        )
        row_fill = PatternFill(start_color="F9FAFB" if target_row % 2 == 0 else "FFFFFF",
                               end_color="F9FAFB" if target_row % 2 == 0 else "FFFFFF",
                               fill_type="solid")

        for col in range(1, ws.max_column + 1):
            c = ws.cell(row=target_row, column=col)
            c.border = thin_border
            if c.fill.fill_type is None:
                c.fill = row_fill

        wb.save(self.excel_path)
        return True

    @staticmethod
    def extract_timestamp_from_path(file_path: str) -> str:
        """สกัด Timestamp จากชื่อไฟล์ เช่น TA200304_20260922-184403.jpg หรือ bangsala-2026-09-30-16-59-08.jpg"""
        base = os.path.basename(file_path)

        m = re.search(r'(\d{4})[-_]?(\d{2})[-_]?(\d{2})[-_](\d{2})[-_]?(\d{2})[-_]?(\d{2})', base)
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)} {m.group(4)}:{m.group(5)}:{m.group(6)}"

        m2 = re.search(r'(\d{4})(\d{2})(\d{2})-(\d{2})(\d{2})(\d{2})', base)
        if m2:
            return f"{m2.group(1)}-{m2.group(2)}-{m2.group(3)} {m2.group(4)}:{m2.group(5)}:{m2.group(6)}"

        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# Singleton instance
excel_logger = WaterLevelExcelLogger()
